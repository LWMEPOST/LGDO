# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import json
import re
import unicodedata
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "generated"
PDF_FONT = Path(r"C:\Windows\Fonts\NotoSansSC-VF.ttf")
if not PDF_FONT.exists():
    PDF_FONT = Path(r"C:\Windows\Fonts\msyh.ttc")

SOURCE_EXTENSIONS = {".md", ".txt", ".csv", ".jsonl"}
SKIP_FILES = {"README.md", "qa_seed_questions.md"}


def iter_source_files() -> list[Path]:
    files: list[Path] = []
    for path in ROOT.rglob("*"):
        if not path.is_file():
            continue
        if "generated" in path.parts:
            continue
        if path.name in SKIP_FILES:
            continue
        if path.suffix.lower() in SOURCE_EXTENSIONS:
            files.append(path)
    return sorted(files)


def safe_output_stem(path: Path) -> str:
    relative = path.relative_to(ROOT).as_posix()
    return relative.replace("/", "__").replace(".", "_")


def set_east_asia_font(run, font_name: str = "Microsoft YaHei") -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9DEE7")
        borders.append(element)
    tbl_pr.append(borders)


def add_docx_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    run.font.size = Pt(10.5)


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        is_separator = all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)
        if not is_separator:
            rows.append(cells)
        index += 1
    return rows, index


def add_table_to_docx(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    set_table_borders(table)
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            text = row[col_index] if col_index < len(row) else ""
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(text)
            set_east_asia_font(run)
            run.font.size = Pt(9.5)
            if row_index == 0:
                run.bold = True
    document.add_paragraph()


def add_text_lines_to_docx(document: Document, lines: list[str], markdown: bool = False) -> None:
    index = 0
    while index < len(lines):
        line = lines[index].rstrip("\n")
        stripped = line.strip()

        if markdown and stripped.startswith("|"):
            rows, next_index = parse_markdown_table(lines, index)
            add_table_to_docx(document, rows)
            index = next_index
            continue

        if not stripped:
            document.add_paragraph()
        elif markdown and stripped.startswith("# "):
            add_docx_paragraph(document, stripped[2:], "Heading 1")
        elif markdown and stripped.startswith("## "):
            add_docx_paragraph(document, stripped[3:], "Heading 2")
        elif markdown and stripped.startswith("### "):
            add_docx_paragraph(document, stripped[4:], "Heading 3")
        elif markdown and stripped.startswith("- "):
            add_docx_paragraph(document, stripped[2:], "List Bullet")
        elif markdown and re.match(r"^\d+\. ", stripped):
            add_docx_paragraph(document, stripped, "List Number")
        elif markdown and stripped.startswith("> "):
            add_docx_paragraph(document, stripped[2:], "Quote")
        else:
            add_docx_paragraph(document, stripped)
        index += 1


def source_as_plain_text(path: Path) -> str:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            rows = list(reader)
        widths = [max(len(row[col]) if col < len(row) else 0 for row in rows) for col in range(max(map(len, rows)))]
        return "\n".join(
            " | ".join((row[col] if col < len(row) else "").ljust(widths[col]) for col in range(len(widths))).rstrip()
            for row in rows
        )

    if path.suffix.lower() == ".jsonl":
        records = []
        with path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if line.strip():
                    records.append(json.loads(line))
        output: list[str] = []
        for record in records:
            output.append(f"ticket_id: {record.get('ticket_id', '')}")
            for key, value in record.items():
                if key != "ticket_id":
                    output.append(f"{key}: {value}")
            output.append("")
        return "\n".join(output).strip()

    return path.read_text(encoding="utf-8")


def create_docx(path: Path, output_path: Path) -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)

    title = safe_output_stem(path)
    document.add_heading(title, level=0)
    add_docx_paragraph(document, f"Source file: {path.relative_to(ROOT).as_posix()}")

    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        add_table_to_docx(document, rows)
    elif suffix == ".jsonl":
        records = []
        with path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if line.strip():
                    records.append(json.loads(line))
        if records:
            keys = list(records[0].keys())
            rows = [keys] + [[str(record.get(key, "")) for key in keys] for record in records]
            add_table_to_docx(document, rows)
    else:
        text = path.read_text(encoding="utf-8")
        add_text_lines_to_docx(document, text.splitlines(), markdown=suffix == ".md")

    document.save(output_path)


def char_width(char: str) -> int:
    return 2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1


def wrap_line(text: str, limit: int = 82) -> list[str]:
    if not text:
        return [""]

    lines: list[str] = []
    current = ""
    width = 0
    for char in text:
        next_width = char_width(char)
        if width + next_width > limit and current:
            lines.append(current)
            current = char
            width = next_width
        else:
            current += char
            width += next_width
    if current:
        lines.append(current)
    return lines


def markdown_to_pdf_lines(text: str) -> list[tuple[str, int]]:
    lines: list[tuple[str, int]] = []
    for raw_line in text.splitlines():
        stripped = raw_line.strip()
        if stripped.startswith("# "):
            content, size = stripped[2:], 17
        elif stripped.startswith("## "):
            content, size = stripped[3:], 14
        elif stripped.startswith("### "):
            content, size = stripped[4:], 12
        elif stripped.startswith("- "):
            content, size = "• " + stripped[2:], 10
        elif re.match(r"^\d+\. ", stripped):
            content, size = stripped, 10
        elif stripped.startswith("> "):
            content, size = stripped[2:], 10
        else:
            content, size = raw_line, 10

        wrapped = wrap_line(content)
        for line in wrapped:
            lines.append((line, size))
    return lines


def create_pdf(path: Path, output_path: Path) -> None:
    source = source_as_plain_text(path)
    if path.suffix.lower() == ".md":
        rendered_lines = markdown_to_pdf_lines(source)
    else:
        rendered_lines = [(line, 10) for raw in source.splitlines() for line in wrap_line(raw)]

    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    font_name = "noto"
    page.insert_font(fontname=font_name, fontfile=str(PDF_FONT))

    margin_x = 48
    y = 56
    bottom = 800

    title = safe_output_stem(path)
    page.insert_text((margin_x, y), title, fontsize=17, fontname=font_name, color=(0.10, 0.16, 0.24))
    y += 24
    source_label = f"Source file: {path.relative_to(ROOT).as_posix()}"
    page.insert_text((margin_x, y), source_label, fontsize=9, fontname=font_name, color=(0.35, 0.39, 0.45))
    y += 24

    for line, size in rendered_lines:
        if y > bottom:
            page = doc.new_page(width=595, height=842)
            page.insert_font(fontname=font_name, fontfile=str(PDF_FONT))
            y = 56

        if not line.strip():
            y += 8
            continue

        color = (0.10, 0.16, 0.24) if size >= 12 else (0.18, 0.20, 0.24)
        page.insert_text((margin_x, y), line, fontsize=size, fontname=font_name, color=color)
        y += max(size + 5, 15)

    doc.subset_fonts()
    doc.save(output_path, garbage=4, deflate=True)
    doc.close()


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    for source in iter_source_files():
        stem = safe_output_stem(source)
        create_docx(source, OUT_DIR / f"{stem}.docx")
        create_pdf(source, OUT_DIR / f"{stem}.pdf")
        print(f"generated: {stem}.docx / {stem}.pdf")


if __name__ == "__main__":
    main()

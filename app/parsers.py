from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


TEXT_EXTENSIONS = {
    ".md",
    ".markdown",
    ".txt",
    ".log",
}

STRUCTURED_TEXT_EXTENSIONS = {
    ".csv",
    ".json",
}

DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
}

LEGACY_BINARY_EXTENSIONS = {
    ".doc",
    ".xls",
    ".ppt",
}

SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | STRUCTURED_TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | LEGACY_BINARY_EXTENSIONS


@dataclass
class ParsedDocument:
    text: str
    parser: str
    warnings: list[str] = field(default_factory=list)
    ocr_used: bool = False
    page_count: int | None = None


def read_document_text(path: Path, ocr_enabled: bool = False) -> ParsedDocument:
    ext = path.suffix.lower()
    if ext in TEXT_EXTENSIONS:
        return ParsedDocument(path.read_text(encoding="utf-8", errors="ignore"), "text")
    if ext == ".csv":
        return ParsedDocument(read_csv(path), "csv")
    if ext == ".json":
        return ParsedDocument(read_json(path), "json")
    if ext == ".pdf":
        return read_pdf(path, ocr_enabled)
    if ext == ".docx":
        return ParsedDocument(read_docx(path), "docx")
    if ext == ".xlsx":
        return ParsedDocument(read_xlsx(path), "xlsx")
    if ext == ".pptx":
        return ParsedDocument(read_pptx(path), "pptx")
    if ext in LEGACY_BINARY_EXTENSIONS:
        return ParsedDocument(
            f"# {path.name}\n\n"
            "该旧版二进制 Office 文件已入库，但第一阶段解析器只直接支持 docx/xlsx/pptx。"
            "建议先另存为新版格式或导出为 Markdown/TXT 后重新扫描。\n",
            f"legacy-{ext.lstrip('.')}",
            ["legacy_office_format"],
        )
    raise ValueError(f"Unsupported file type: {path.suffix}")


def read_csv(path: Path) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as file:
        reader = csv.reader(file)
        for index, row in enumerate(reader):
            if index >= 200:
                rows.append("... CSV 内容超过 200 行，已截断 ...")
                break
            rows.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(rows)


def read_json(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return raw
    return json.dumps(parsed, ensure_ascii=False, indent=2)


def read_pdf(path: Path, ocr_enabled: bool = False) -> ParsedDocument:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("缺少 pymupdf 依赖，无法解析 PDF。请运行 pip install -e .") from exc

    parts: list[str] = []
    with fitz.open(str(path)) as document:
        page_count = document.page_count
        for index, page in enumerate(document, start=1):
            text = page.get_text("text") or ""
            if text.strip():
                parts.append(f"## Page {index}\n\n{text.strip()}")
    if parts:
        return ParsedDocument("\n\n".join(parts), "pdf-text", page_count=page_count)
    warning = "pdf_text_empty"
    if ocr_enabled:
        return ParsedDocument(
            f"# {path.name}\n\nOCR 通道已启用，但当前 MVP 未配置 OCR 引擎。请接入 PaddleOCR/Tesseract 后重试。",
            "pdf-ocr-placeholder",
            [warning, "ocr_engine_not_configured"],
            ocr_used=True,
            page_count=page_count,
        )
    return ParsedDocument(
        f"# {path.name}\n\n未能从 PDF 中提取文本，可能需要 OCR。",
        "pdf-text",
        [warning, "ocr_disabled"],
        page_count=page_count,
    )


def read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 依赖，无法解析 DOCX。请运行 pip install -e .") from exc

    document = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl 依赖，无法解析 XLSX。请运行 pip install -e .") from exc

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_index > 200:
                parts.append("... Sheet 内容超过 200 行，已截断 ...")
                break
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                parts.append(" | ".join(values))
    return "\n".join(parts)


def read_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("缺少 python-pptx 依赖，无法解析 PPTX。请运行 pip install -e .") from exc

    presentation = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = collect_slide_text(slide)
        notes_text = collect_notes_text(slide)
        if notes_text:
            texts.append("Notes:\n" + notes_text)
        if texts:
            parts.append(f"# Slide {index}\n" + "\n".join(texts))
    return "\n\n".join(parts)


def collect_slide_text(slide: Any) -> list[str]:
    texts: list[str] = []
    for shape in slide.shapes:
        texts.extend(collect_shape_text(shape))
    return dedupe_preserve_order(texts)


def collect_shape_text(shape: Any) -> list[str]:
    parts: list[str] = []

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [normalize_shape_text(cell.text) for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    if getattr(shape, "has_text_frame", False):
        text = normalize_shape_text(shape.text)
        if text:
            parts.append(text)

    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            parts.extend(collect_shape_text(child))

    return parts


def collect_notes_text(slide: Any) -> str:
    try:
        notes_slide = slide.notes_slide
    except Exception:
        return ""
    texts: list[str] = []
    for shape in notes_slide.shapes:
        texts.extend(collect_shape_text(shape))
    return "\n".join(dedupe_preserve_order(texts))


def normalize_shape_text(text: str) -> str:
    lines = [line.strip() for line in (text or "").splitlines()]
    return "\n".join(line for line in lines if line)


def dedupe_preserve_order(items: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        value = item.strip()
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result

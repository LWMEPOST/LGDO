from pptx import Presentation
from pptx.util import Inches

from app.normalization import chunk_markdown
from app.parsers import read_pdf, read_pptx


def test_pptx_parser_extracts_slide_tables(tmp_path):
    pptx_path = tmp_path / "roadmap.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "2025 H2 (进行中)"
    table_shape = slide.shapes.add_table(3, 4, Inches(1), Inches(1.5), Inches(8), Inches(2))
    rows = [
        ["功能", "优先级", "目标日期", "进度"],
        ["中文文字渲染优化", "P0", "2025-08", "60%"],
        ["实时协作画布", "P0", "2025-09", "30%"],
    ]
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table_shape.table.cell(row_index, col_index).text = value
    presentation.save(pptx_path)

    text = read_pptx(pptx_path)

    assert "2025 H2" in text
    assert "功能 | 优先级 | 目标日期 | 进度" in text
    assert "中文文字渲染优化 | P0 | 2025-08 | 60%" in text


def test_pdf_parser_uses_fitz_text_layer(tmp_path):
    import fitz

    pdf_path = tmp_path / "policy.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "standard work time 9:00-18:00")
    document.save(pdf_path)
    document.close()

    parsed = read_pdf(pdf_path)

    assert parsed.parser == "pdf-text"
    assert parsed.page_count == 1
    assert "standard work time" in parsed.text


def test_chunk_markdown_carries_overlap_between_chunks():
    markdown = "# Policy\n\n" + "\n\n".join(
        [
            "第一段 " + "A" * 180,
            "第二段包含婚假 10天 直属上级 HR",
            "第三段 " + "B" * 180,
            "第四段 " + "C" * 180,
        ]
    )

    chunks = chunk_markdown("src_test", markdown, {"title": "Policy", "domain": "customer_service"}, max_chars=300, overlap_chars=80)

    assert len(chunks) > 1
    assert any("第二段包含婚假" in chunk["text"] for chunk in chunks)
    assert any(
        "第二段包含婚假" in chunks[index]["text"] and "第三段" in chunks[index]["text"]
        for index in range(1, len(chunks))
    )
